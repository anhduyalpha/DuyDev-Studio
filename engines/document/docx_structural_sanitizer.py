#!/usr/bin/env python3
"""
DD Studio - DOCX Structural Sanitizer (< 230 lines)
Post-conversion safety net using python-docx to unpack collapsed tables,
clean header boxes, restore scientific exponents, and normalize question paragraphs.
"""

import os
import re
from typing import Dict, Any, List
import docx
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

EXPONENT_RE = re.compile(
    r'\b(10|2)(1[0-9]|2[0-9]|3[0-9]|4[0-9]|5[0-9]|6[0-9])\s*(bytes?|B|bit|bits|Hz|kHz|MHz|GHz|s|ms|μs|us|ns)\b'
)
QUESTION_SPLIT_RE = re.compile(
    r'(?:^|[\r\n]+|[.!?]\s+)(\d+\.\d+|\d+\.)\s*(?=[\t\r\n]|\s{2,}|[A-ZÀ-Ỹ])',
    re.UNICODE
)
QUESTION_TOKEN_RE = re.compile(r'(\d+\.\d+|\d+\.)')
QUESTION_PREFIX_RE = re.compile(r'^(\d+\.\d+|\d+\.)\s*[\t\s]+')
HEADER_KW_RE = re.compile(
    r'(?:khoa|trường|đại học|học viện|bộ môn|viện|trang|page|faculty|university|department|semester|\b(?:19|20)\d{2}\b)',
    re.IGNORECASE
)


def _format_runs_with_exponents(p: Paragraph, text: str, restore_exp: bool = True):
    q_match = QUESTION_PREFIX_RE.match(text)
    if q_match:
        q_num = q_match.group(1)
        r_num = p.add_run(f"{q_num}\t")
        r_num.bold = True
        text = text[q_match.end():]

    if not restore_exp or not EXPONENT_RE.search(text):
        p.add_run(text)
        return

    last_idx = 0
    for m in EXPONENT_RE.finditer(text):
        if m.start() > last_idx:
            p.add_run(text[last_idx:m.start()])
        base, exp, unit = m.group(1), m.group(2), m.group(3)
        p.add_run(base)
        exp_run = p.add_run(exp)
        exp_run.font.superscript = True
        p.add_run(" " + unit)
        last_idx = m.end()
    if last_idx < len(text):
        p.add_run(text[last_idx:])


def _clean_header_tables(doc: docx.Document) -> int:
    removed = 0
    tables_to_clean = []
    for tbl in doc.tables:
        if len(tbl.rows) == 1 and len(tbl.columns) == 2:
            t0 = tbl.cell(0, 0).text.strip()
            t1 = tbl.cell(0, 1).text.strip()
            full = f"{t0} {t1}"
            if len(full) < 160 and HEADER_KW_RE.search(full):
                tables_to_clean.append((tbl, t0, t1))

    for tbl, t0, t1 in tables_to_clean:
        p_el = OxmlElement('w:p')
        tbl._element.addprevious(p_el)
        p = Paragraph(p_el, tbl._element.getparent())
        p.add_run(f"{t0}\t{t1}")
        tbl._element.getparent().remove(tbl._element)
        removed += 1
    return removed


def _unpack_collapsed_tables(doc: docx.Document, restore_exp: bool) -> int:
    unpacked = 0
    tables_to_unpack = []
    for tbl in doc.tables:
        if len(tbl.rows) == 1 and len(tbl.columns) == 2:
            c0_text = tbl.cell(0, 0).text.strip()
            nums = QUESTION_TOKEN_RE.findall(c0_text)
            if len(nums) >= 1:
                tables_to_unpack.append((tbl, nums))

    for tbl, nums in tables_to_unpack:
        c1 = tbl.cell(0, 1)
        raw_lines = [s.strip() for p in c1.paragraphs for s in p.text.split("\n") if s.strip()]
        parent = tbl._element.getparent()

        for idx, q_num in enumerate(nums):
            desc = raw_lines[idx] if idx < len(raw_lines) else ""
            p_el = OxmlElement('w:p')
            tbl._element.addprevious(p_el)
            p = Paragraph(p_el, parent)
            num_run = p.add_run(f"{q_num}\t")
            num_run.bold = True
            _format_runs_with_exponents(p, desc, restore_exp)

        if len(raw_lines) > len(nums):
            for extra in raw_lines[len(nums):]:
                p_el = OxmlElement('w:p')
                tbl._element.addprevious(p_el)
                p = Paragraph(p_el, parent)
                _format_runs_with_exponents(p, extra, restore_exp)

        parent.remove(tbl._element)
        unpacked += 1
    return unpacked


def _normalize_question_paragraphs(doc: docx.Document, restore_exp: bool) -> int:
    normalized = 0
    for p in list(doc.paragraphs):
        raw_text = p.text
        cleaned = raw_text.strip()
        if cleaned.lower().startswith(('bài ', 'bài\t', 'câu ', 'câu\t')) and len(cleaned) < 40:
            continue

        matches = list(QUESTION_SPLIT_RE.finditer(raw_text))
        if len(matches) > 1:
            parent = p._element.getparent()
            if matches[0].start() > 0:
                prefix = raw_text[:matches[0].start()].strip()
                if prefix:
                    p_el = OxmlElement('w:p')
                    p._element.addprevious(p_el)
                    p_pre = Paragraph(p_el, parent)
                    _format_runs_with_exponents(p_pre, prefix, restore_exp)

            for i, m in enumerate(matches):
                q_num = m.group(1)
                start_pos = m.end()
                end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(raw_text)
                q_body = raw_text[start_pos:end_pos].strip()
                p_el = OxmlElement('w:p')
                p._element.addprevious(p_el)
                new_p = Paragraph(p_el, parent)
                r_num = new_p.add_run(f"{q_num}\t")
                r_num.bold = True
                _format_runs_with_exponents(new_p, q_body, restore_exp)
                normalized += 1

            parent.remove(p._element)
        elif len(matches) == 1 and matches[0].start() == 0:
            q_num = matches[0].group(1)
            for r in p.runs:
                if r.text.strip():
                    if r.text.strip().startswith(q_num):
                        r.bold = True
                    break
    return normalized


def _restore_remaining_exponents(doc: docx.Document) -> int:
    restored = 0
    paragraphs_to_check: List[Paragraph] = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                paragraphs_to_check.extend(cell.paragraphs)

    for p in paragraphs_to_check:
        if EXPONENT_RE.search(p.text):
            has_superscript = any(r.font.superscript for r in p.runs)
            if not has_superscript:
                old_text = p.text
                p.text = ""
                _format_runs_with_exponents(p, old_text, restore_exp=True)
                restored += 1
    return restored


def sanitize_docx(input_docx: str, output_docx: str, rules: Dict[str, Any]) -> Dict[str, Any]:
    """Sanitizes converted DOCX against common pdf2docx structural bugs."""
    if not os.path.isfile(input_docx):
        raise FileNotFoundError(f"Input DOCX not found: {input_docx}")

    doc = docx.Document(input_docx)
    stats = {
        "headers_cleaned": 0,
        "tables_unpacked": 0,
        "questions_normalized": 0,
        "exponents_restored": 0
    }

    if rules.get("clean_header_footer_tables", True):
        stats["headers_cleaned"] = _clean_header_tables(doc)

    restore_exp = rules.get("restore_scientific_exponents", True)

    if rules.get("unpack_collapsed_list_tables", True):
        stats["tables_unpacked"] = _unpack_collapsed_tables(doc, restore_exp)

    if rules.get("normalize_question_paragraphs", True):
        stats["questions_normalized"] = _normalize_question_paragraphs(doc, restore_exp)

    if restore_exp:
        stats["exponents_restored"] = _restore_remaining_exponents(doc)

    doc.save(output_docx)
    return stats


if __name__ == "__main__":
    import sys
    inp = sys.argv[1] if len(sys.argv) > 1 else ""
    out = sys.argv[2] if len(sys.argv) > 2 else "sanitized.docx"
    if inp and os.path.exists(inp):
        res = sanitize_docx(inp, out, {
            "clean_header_footer_tables": True,
            "unpack_collapsed_list_tables": True,
            "normalize_question_paragraphs": True,
            "restore_scientific_exponents": True
        })
        print("Sanitizer result:", res)

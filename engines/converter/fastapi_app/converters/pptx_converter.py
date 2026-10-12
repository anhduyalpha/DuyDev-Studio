"""
PPTX Converter - Enhanced Version 2.0
Converts PowerPoint files to PDF, TXT, and multi-slide images (PNG, JPG, WEBP).
Uses PyMuPDF for slide image rendering.
Multi-slide images automatically bundled into ZIP archives.
"""
import os
import asyncio
import subprocess
import shutil
import zipfile


async def convert_pptx(input_path: str, output_dir: str, target_format: str) -> dict:
    """PPTX converter supporting PDF, TXT, and slide image outputs."""
    loop = asyncio.get_event_loop()
    return await loop.run_in_executor(None, _process_pptx, input_path, output_dir, target_format)


def _process_pptx(input_path: str, output_dir: str, target_format: str) -> dict:
    try:
        filename = os.path.basename(input_path)
        name, _ = os.path.splitext(filename)
        target_fmt = target_format.lower().lstrip('.')
        output_filename = f"{name}.{target_fmt}"
        output_path = os.path.join(output_dir, output_filename)

        if target_fmt == 'pdf':
            return _pptx_to_pdf(input_path, output_path, output_filename)
        elif target_fmt in ['png', 'jpg', 'jpeg', 'webp']:
            return _pptx_to_images(input_path, output_dir, target_fmt, name)
        elif target_fmt == 'txt':
            return _pptx_to_txt(input_path, output_path, output_filename)
        else:
            return {"success": False, "error": f"Định dạng đích không được hỗ trợ cho PPTX: {target_fmt}"}

    except Exception as e:
        return {"success": False, "error": f"Lỗi chuyển đổi PPTX: {str(e)}"}


def _pptx_to_pdf(input_path: str, output_path: str, output_filename: str) -> dict:
    """Convert PPTX to PDF using LibreOffice or PowerPoint COM."""
    # Method 1: Try LibreOffice (cross-platform)
    try:
        libreoffice_paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "/usr/bin/soffice",
            "/usr/bin/libreoffice",
            "/usr/local/bin/soffice",
            "soffice",
            "libreoffice",
            "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        ]

        soffice_path = None
        for p in libreoffice_paths:
            if os.path.exists(p):
                soffice_path = p
                break
            resolved = shutil.which(p)
            if resolved and os.path.exists(resolved):
                soffice_path = resolved
                break

        if soffice_path:
            out_dir = os.path.dirname(output_path)
            subprocess.run([
                soffice_path,
                "--headless",
                "--convert-to", "pdf",
                "--outdir", out_dir,
                input_path
            ], capture_output=True, text=True, timeout=180)

            expected_output = os.path.join(out_dir, os.path.splitext(os.path.basename(input_path))[0] + ".pdf")
            if os.path.exists(expected_output):
                if expected_output != output_path:
                    shutil.move(expected_output, output_path)
                return {"success": True, "output_path": output_path, "filename": output_filename}
    except Exception as e:
        print(f"[PPTX→PDF] LibreOffice failed: {e}")

    # Method 2: Try PowerPoint COM (Windows only)
    try:
        import platform
        if platform.system() == 'Windows':
            import comtypes.client

            powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
            powerpoint.Visible = 1

            presentation = powerpoint.Presentations.Open(os.path.abspath(input_path))
            presentation.SaveAs(os.path.abspath(output_path), 32)  # 32 = ppSaveAsPDF
            presentation.Close()
            powerpoint.Quit()

            if os.path.exists(output_path):
                return {"success": True, "output_path": output_path, "filename": output_filename}
    except Exception as e:
        print(f"[PPTX→PDF] PowerPoint COM failed: {e}")

    return {
        "success": False,
        "error": "Chuyển đổi PPTX sang PDF yêu cầu LibreOffice hoặc Microsoft PowerPoint trên máy chủ."
    }


def _pptx_to_images(input_path: str, output_dir: str, target_format: str, name: str) -> dict:
    """Convert PPTX to images using PyMuPDF. Auto-bundles multi-slide into ZIP archive."""
    target_fmt = target_format.lower().lstrip('.')

    # First convert to PDF
    pdf_path = os.path.join(output_dir, f"{name}_temp.pdf")
    pdf_result = _pptx_to_pdf(input_path, pdf_path, f"{name}.pdf")

    if not pdf_result["success"]:
        # Fallback: Extract embedded images from PPTX directly
        try:
            images_extracted = []
            with zipfile.ZipFile(input_path, 'r') as pptx:
                for i, item in enumerate(pptx.namelist()):
                    if item.startswith('ppt/media/') and any(item.lower().endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.gif', '.webp']):
                        img_data = pptx.read(item)
                        original_ext = os.path.splitext(item)[1]
                        img_filename = f"{name}_anh_{i+1}{original_ext}"
                        img_path = os.path.join(output_dir, img_filename)
                        with open(img_path, 'wb') as f:
                            f.write(img_data)
                        images_extracted.append(img_filename)

            if images_extracted:
                if len(images_extracted) == 1:
                    single_path = os.path.join(output_dir, images_extracted[0])
                    return {"success": True, "output_path": single_path, "filename": images_extracted[0], "all_files": images_extracted}

                # Bundle into ZIP
                zip_filename = f"{name}_anh.zip"
                zip_path = os.path.join(output_dir, zip_filename)
                with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
                    for f in images_extracted:
                        f_path = os.path.join(output_dir, f)
                        zf.write(f_path, arcname=f)

                return {
                    "success": True,
                    "output_path": zip_path,
                    "filename": zip_filename,
                    "all_files": images_extracted,
                    "is_zip": True,
                    "note": f"Đã trích xuất {len(images_extracted)} hình ảnh từ bài thuyết trình"
                }

            return {"success": False, "error": "Không tìm thấy hình ảnh nào trong bài thuyết trình và chuyển đổi PDF thất bại"}
        except Exception as e:
            return {"success": False, "error": f"Lỗi trích xuất hình ảnh: {str(e)}"}

    # Convert PDF slides to images using PyMuPDF
    try:
        try:
            import pymupdf
        except ImportError:
            import fitz as pymupdf
    except ImportError:
        if os.path.exists(pdf_path):
            os.remove(pdf_path)
        return {"success": False, "error": "Thư viện PyMuPDF chưa được cài đặt. Vui lòng chạy 'pip install PyMuPDF'"}

    try:
        pdf_doc = pymupdf.open(pdf_path)
        output_files = []

        for page_num in range(len(pdf_doc)):
            page = pdf_doc[page_num]
            mat = pymupdf.Matrix(2.0, 2.0)
            pix = page.get_pixmap(matrix=mat)

            img_filename = f"{name}_slide_{page_num + 1:02d}.{target_fmt}"
            img_path = os.path.join(output_dir, img_filename)

            if target_fmt in ['jpg', 'jpeg']:
                pix.save(img_path, output="jpeg", jpg_quality=95)
            elif target_fmt == 'webp':
                from PIL import Image
                mode = "RGB" if pix.n < 4 else "RGBA"
                pil_img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
                pil_img.save(img_path, 'WEBP', quality=95)
            else:
                pix.save(img_path)

            output_files.append(img_filename)

        pdf_doc.close()

        if os.path.exists(pdf_path):
            os.remove(pdf_path)

        if not output_files:
            return {"success": False, "error": "Không thể kết xuất trang slide thành hình ảnh"}

        # Single slide: return image directly
        if len(output_files) == 1:
            single_path = os.path.join(output_dir, output_files[0])
            return {
                "success": True,
                "output_path": single_path,
                "filename": output_files[0],
                "all_files": output_files
            }

        # Multi-slide: automatically bundle into ZIP
        zip_filename = f"{name}_{target_fmt}.zip"
        zip_path = os.path.join(output_dir, zip_filename)
        with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
            for f in output_files:
                f_path = os.path.join(output_dir, f)
                zf.write(f_path, arcname=f)

        return {
            "success": True,
            "output_path": zip_path,
            "filename": zip_filename,
            "all_files": output_files,
            "is_zip": True,
            "note": f"Đã xuất toàn bộ {len(output_files)} slide và đóng gói vào tệp zip"
        }

    except Exception as e:
        if os.path.exists(pdf_path):
            os.remove(pdf_path)
        return {"success": False, "error": f"Lỗi kết xuất ảnh từ slide PPTX: {str(e)}"}


def _pptx_to_txt(input_path: str, output_path: str, output_filename: str) -> dict:
    """Extract text from PPTX."""
    try:
        from pptx import Presentation
    except ImportError:
        return {"success": False, "error": "Thư viện python-pptx chưa được cài đặt. Vui lòng chạy 'pip install python-pptx'"}

    try:
        prs = Presentation(input_path)
        text_content = []

        for slide_num, slide in enumerate(prs.slides, 1):
            text_content.append(f"=== Slide {slide_num} ===\n")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_content.append(shape.text)
            text_content.append("\n")

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write("\n".join(text_content))

        return {"success": True, "output_path": output_path, "filename": output_filename}
    except Exception as e:
        return {"success": False, "error": f"Lỗi trích xuất văn bản từ PPTX: {str(e)}"}
